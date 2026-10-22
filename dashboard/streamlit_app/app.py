
import io

import requests
import streamlit as st
from docx import Document
from pypdf import PdfReader


# ==================================================
# CONFIGURATION
# ==================================================

st.set_page_config(
    page_title="CareerLens AI",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
)

API_URL = "http://127.0.0.1:8000/career/report"


# ==================================================
# CUSTOM STYLING
# ==================================================

st.markdown(
    """
    <style>
    .block-container {
        max-width: 1250px;
        padding-top: 2rem;
        padding-bottom: 3rem;
    }

    .hero {
        padding: 28px;
        border-radius: 18px;
        background: linear-gradient(120deg, #172554, #312e81, #4c1d95);
        color: white;
        margin-bottom: 25px;
    }

    .hero h1 {
        color: white;
        font-size: 38px;
        margin-bottom: 8px;
    }

    .hero p {
        color: #e0e7ff;
        font-size: 16px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ==================================================
# RESUME TEXT EXTRACTION (FOR PREVIEW)
# ==================================================

def extract_resume_text(uploaded_file):
    """Extract text from a PDF or DOCX file."""

    file_bytes = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))

        if reader.is_encrypted:
            raise ValueError(
                "This PDF is password-protected. Upload an unlocked PDF."
            )

        extracted_pages = []

        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                extracted_pages.append(page_text)

        text = "\n".join(extracted_pages)

    elif filename.endswith(".docx"):
        document = Document(io.BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        table_lines = []

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    table_lines.append(" | ".join(cells))

        text = "\n".join(paragraphs + table_lines)

    else:
        raise ValueError("Please upload a PDF or DOCX file.")

    if not text.strip():
        raise ValueError(
            "No readable text found. If your PDF is scanned, "
            "it may require OCR."
        )

    return text.strip()


# ==================================================
# DISPLAY CAREER REPORT
# ==================================================

def display_report(report):

    st.success("Resume analysis completed!")

    st.header("📊 Your Career Overview")

    score = report.get("resume_score", 0)
    recommendations = report.get("recommendations") or []
    strengths = report.get("strengths") or []
    weaknesses = report.get("weaknesses") or []
    roadmap = report.get("roadmap") or []

    col1, col2, col3 = st.columns(3)

    with col1:
        st.metric("Resume Score", f"{score}/100")

    with col2:
        st.metric("Job Recommendations", len(recommendations))

    with col3:
        st.metric("Improvement Areas", len(weaknesses))

    try:
        st.progress(
            max(0.0, min(float(score) / 100, 1.0)),
            text=f"Resume Score: {score}/100",
        )
    except (TypeError, ValueError):
        pass

    career_level = report.get("career_level")
    if career_level:
        st.subheader("🎯 Career Level")
        st.info(career_level)

    career_summary = report.get("career_summary")
    if career_summary:
        st.write(career_summary)

    # Strengths and weaknesses
    st.divider()
    left, right = st.columns(2)

    with left:
        st.subheader("💪 Your Strengths")
        if strengths:
            for item in strengths:
                st.write(f"✅ {item}")
        else:
            st.write("No strengths returned.")

    with right:
        st.subheader("📈 Skills to Improve")
        if weaknesses:
            for item in weaknesses:
                st.write(f"• {item}")
        else:
            st.write("No improvement areas returned.")

    # Job recommendations
    st.divider()
    st.header("💼 Recommended Jobs")

    if recommendations:
        for index, job in enumerate(recommendations, start=1):

            if not isinstance(job, dict):
                st.write(job)
                continue

            title = job.get("title", "Job Opportunity")
            company = job.get("company", "Company not specified")
            location = job.get("location", "Not specified")
            match_score = job.get(
                "match_score",
                job.get("similarity_score"),
            )

            heading = f"{index}. {title} — {company}"

            if match_score is not None:
                heading += f" | Match: {match_score}"

            with st.expander(heading):

                st.write(f"**Location:** {location}")

                if match_score is not None:
                    st.write(f"**Match Score:** {match_score}")

                matched = job.get("matched_skills") or []
                missing = job.get("missing_skills") or []

                st.markdown("**Matched Skills**")
                st.write(
                    ", ".join(map(str, matched))
                    if matched else "Not available"
                )

                st.markdown("**Skills to Develop**")
                st.write(
                    ", ".join(map(str, missing))
                    if missing else "Not available"
                )

                description = job.get("description")
                if description:
                    st.write(description)

                link = job.get("link") or job.get("url")

                if isinstance(link, str) and link.startswith(
                    ("https://", "http://")
                ):
                    st.markdown(f"[View Job Posting]({link})")

    else:
        st.info("No job recommendations were returned.")

    # Learning roadmap
    st.divider()
    st.header("🗓️ Your Learning Roadmap")

    if roadmap:
        for index, week in enumerate(roadmap, start=1):

            if not isinstance(week, dict):
                st.write(week)
                continue

            week_number = week.get("week", index)
            focus = week.get("focus", "Learning Plan")

            with st.expander(
                f"Week {week_number}: {focus}",
                expanded=True,
            ):
                goal = week.get("goal", "")
                practice = week.get("practice", "")

                if goal:
                    st.markdown("**Goal**")
                    st.write(goal)

                if practice:
                    st.markdown("**Practice**")
                    st.write(practice)

                topics = week.get("topics") or []
                if topics:
                    st.markdown("**Topics to Learn**")
                    for topic in topics:
                        st.write(f"- {topic}")

                tasks = week.get("tasks") or []
                if tasks:
                    st.markdown("**Tasks**")
                    for task in tasks:
                        st.write(f"- {task}")

    else:
        st.info("No learning roadmap was returned.")

    with st.expander("View Technical Report"):
        st.json(report)


# ==================================================
# SIDEBAR
# ==================================================

with st.sidebar:

    st.title("🎯 CareerLens AI")
    st.caption("AI-Powered Career Intelligence")

    st.divider()

    page = st.radio(
        "Navigation",
        ["Resume Analyzer", "About CareerLens AI"],
    )

    st.divider()

    st.markdown("### Platform Features")
    st.markdown(
        """
        - 📄 Resume Analysis
        - 📊 Resume Scoring
        - 💪 Strength Analysis
        - 📈 Skill Gap Detection
        - 💼 Job Recommendations
        - 🗓️ Learning Roadmap
        """
    )


# ==================================================
# RESUME ANALYZER PAGE
# ==================================================

if page == "Resume Analyzer":

    st.markdown(
        """
        <div class="hero">
            <h1>🎯 CareerLens AI</h1>
            <p>
                Turn your resume into a personalized career plan.
                Discover job opportunities, identify skill gaps,
                and plan your next career move.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.header("📄 Upload Your Resume")

    st.write(
        "Choose your resume as a PDF or Word document. "
        "Your file will be sent directly to the CareerLens AI backend."
    )

    uploaded_resume = st.file_uploader(
        "Choose your resume",
        type=["pdf", "docx"],
        help="Supported formats: PDF and DOCX.",
    )

    if uploaded_resume:

        st.success(f"Selected file: {uploaded_resume.name}")

        st.caption(
            f"File size: {uploaded_resume.size / 1024:.1f} KB"
        )

        # Preview extracted text
        with st.expander("Preview extracted resume text"):

            try:
                preview_text = extract_resume_text(uploaded_resume)

                st.text_area(
                    "Extracted text",
                    value=preview_text[:10000],
                    height=200,
                    disabled=True,
                )

                if len(preview_text) > 10000:
                    st.caption(
                        "Preview shortened. Your original file "
                        "will be sent for analysis."
                    )

            except Exception as error:
                st.warning(f"Preview unavailable: {error}")

        # Analyze button
        analyze = st.button(
            "🚀 Analyze My Resume",
            type="primary",
            use_container_width=True,
        )

        if analyze:

            try:
                with st.spinner(
                    "Analyzing your resume and finding matching jobs..."
                ):

                    # IMPORTANT:
                    # The backend expects a multipart file field
                    # named "file", not JSON containing resume_text.
                    response = requests.post(
                        API_URL,
                        files={
                            "file": (
                                uploaded_resume.name,
                                uploaded_resume.getvalue(),
                                uploaded_resume.type
                                or "application/octet-stream",
                            )
                        },
                        timeout=180,
                    )

                    if not response.ok:
                        st.error(
                            f"Backend returned HTTP "
                            f"{response.status_code}."
                        )
                        st.code(response.text[:3000])
                        st.stop()

                    report = response.json()

                st.divider()
                display_report(report)

            except requests.exceptions.ConnectionError:
                st.error(
                    "Cannot connect to the backend. "
                    "Make sure FastAPI is running at "
                    "http://127.0.0.1:8000."
                )

            except requests.exceptions.Timeout:
                st.error(
                    "The request timed out. Check the FastAPI terminal."
                )

            except requests.exceptions.RequestException as error:
                st.error(f"API request failed: {error}")

            except Exception as error:
                st.error(f"Unable to analyze resume: {error}")

    else:
        st.info("👆 Upload a PDF or DOCX resume to get started.")

        st.divider()
        st.subheader("What you will discover")

        col1, col2, col3 = st.columns(3)

        with col1:
            st.markdown("### 📊 Resume Score")
            st.write("Understand your resume's current score.")

        with col2:
            st.markdown("### 💼 Job Matches")
            st.write("Discover relevant job opportunities.")

        with col3:
            st.markdown("### 🗓️ Learning Roadmap")
            st.write("Identify skills to improve over four weeks.")


# ==================================================
# ABOUT PAGE
# ==================================================

elif page == "About CareerLens AI":

    st.title("About CareerLens AI")

    st.write(
        """
        CareerLens AI is a career intelligence platform that
        connects resume analysis with job-market insights.

        It helps users:
        - Analyze resume scores and strengths.
        - Discover job recommendations.
        - Identify missing skills.
        - Follow a structured learning roadmap.
        """
    )

    st.caption(
        "Built with Python, Streamlit, FastAPI and PostgreSQL."
    )
